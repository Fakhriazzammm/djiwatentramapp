"""
Yumi GM Agent — Schedule Export Module

Generates Excel (.xlsx) roster files from Supabase schedule data.
"""

import io
import logging
from datetime import datetime, timedelta, timezone
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logger = logging.getLogger("yumi-ai-bot")

WIB = timezone(timedelta(hours=7))

# ─── Style constants ────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1a5dc6", end_color="1a5dc6", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
SUBHEADER_FILL = PatternFill(start_color="e8edf5", end_color="e8edf5", fill_type="solid")
ROLE_FONT = Font(name="Calibri", size=9, italic=True, color="666666")
SHIFT_FONT = Font(name="Calibri", size=10, bold=True)

SHIFT_FILLS = {
    "PAGI": PatternFill(start_color="d1fae5", end_color="d1fae5", fill_type="solid"),
    "SORE": PatternFill(start_color="e0f2fe", end_color="e0f2fe", fill_type="solid"),
    "GENERAL": PatternFill(start_color="f1f5f9", end_color="f1f5f9", fill_type="solid"),
}

OFF_FILL = PatternFill(start_color="fee2e2", end_color="fee2e2", fill_type="solid")
OFF_FONT = Font(name="Calibri", size=10, color="dc2626")

THIN_BORDER = Border(
    left=Side(style="thin", color="d0d0d0"),
    right=Side(style="thin", color="d0d0d0"),
    top=Side(style="thin", color="d0d0d0"),
    bottom=Side(style="thin", color="d0d0d0"),
)


def generate_roster_excel(
    schedules: list[dict],
    month: int,
    year: int,
    staff_columns: list[str] | None = None,
) -> bytes:
    """
    Generate an Excel roster from schedule data.

    Args:
        schedules: List of dicts with keys: userName, userSlug, roleName,
                   date, shiftName, startTime, endTime, status
        month: Month number (1-12)
        year: Year as 4-digit
        staff_columns: Ordered list of staff names (optional, auto-detected)

    Returns:
        Excel file bytes
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{_month_name(month).upper()} {year}"

    # Collect unique staff in order
    if not staff_columns:
        seen = set()
        staff_order = []
        for s in schedules:
            name = s.get("userName", "").strip()
            if name and name not in seen:
                seen.add(name)
                staff_order.append(name)
        staff_columns = staff_order

    # Collect unique dates
    dates = sorted(set(
        s.get("date", "") for s in schedules if s.get("date")
    ))

    # Build lookup: {date_staffName: schedule}
    lookup: dict[str, dict] = {}
    for s in schedules:
        key = f"{s.get('date')}_{s.get('userName', '').strip()}"
        lookup[key] = s

    # ── Build rows ──
    rows_data = []
    week_num = 1
    for i, date_str in enumerate(dates):
        dt = datetime.strptime(date_str, "%Y-%m-%d").replace(tzinfo=WIB)
        day_name = _day_name_ina(dt.weekday())

        # Week boundary
        if i > 0 and dt.weekday() == 0:
            week_num += 1

        row = {
            "week": f"Minggu Ke {week_num}",
            "day": day_name,
            "date": dt.strftime("%d %b"),
            "date_sort": date_str,
            "cells": [],
        }
        for staff_name in staff_columns:
            key = f"{date_str}_{staff_name}"
            sch = lookup.get(key)
            row["cells"].append(sch)
        rows_data.append(row)

    # ── Write Excel ──

    # Row 1: Title
    ws.merge_cells(
        start_row=1, start_column=1,
        end_row=1, end_column=3 + len(staff_columns)
    )
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"JADWAL {_month_name(month).upper()} {year}"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="1a5dc6")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Row 3: Headers
    headers = ["MINGGU KE", "HARI", "TANGGAL"]
    headers += list(staff_columns)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    ws.row_dimensions[3].height = 22

    # Data rows (starting row 4)
    for ri, row_data in enumerate(rows_data, 4):
        # Week column
        c_week = ws.cell(row=ri, column=1, value=row_data["week"])
        c_week.font = Font(name="Calibri", size=9, color="666666")
        c_week.alignment = Alignment(horizontal="center", vertical="center")
        c_week.border = THIN_BORDER

        # Day column
        c_day = ws.cell(row=ri, column=2, value=row_data["day"])
        c_day.font = Font(name="Calibri", size=10, bold=True)
        c_day.alignment = Alignment(horizontal="center", vertical="center")
        c_day.border = THIN_BORDER

        # Date column
        c_date = ws.cell(row=ri, column=3, value=row_data["date"])
        c_date.font = Font(name="Calibri", size=10)
        c_date.alignment = Alignment(horizontal="center", vertical="center")
        c_date.border = THIN_BORDER

        # Staff cells
        for ci, sch in enumerate(row_data["cells"]):
            col = 4 + ci
            cell = ws.cell(row=ri, column=col)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

            if sch:
                shift_name = (sch.get("shiftName") or "").upper()
                status = (sch.get("status") or "").lower()
                start_time = sch.get("startTime") or ""
                end_time = sch.get("endTime") or ""

                if status == "off":
                    cell.value = "OFF"
                    cell.fill = OFF_FILL
                    cell.font = OFF_FONT
                elif shift_name == "GENERAL":
                    cell.value = "GENERAL"
                    cell.fill = SHIFT_FILLS.get("GENERAL")
                    cell.font = Font(name="Calibri", size=10, color="475569")
                else:
                    time_str = ""
                    if start_time:
                        # Truncate seconds
                        st = start_time[:5] if len(start_time) >= 5 else start_time
                        et = end_time[:5] if end_time and len(end_time) >= 5 else ""
                        time_str = f" ({st} - {et})" if et else f" ({st})"
                    cell.value = f"{shift_name}{time_str}"
                    fill = SHIFT_FILLS.get(shift_name)
                    if fill:
                        cell.fill = fill
                    cell.font = Font(name="Calibri", size=10, bold=True)
            else:
                cell.value = "-"
                cell.font = Font(name="Calibri", size=9, color="cccccc")

        # Row height
        ws.row_dimensions[ri].height = 20

    # Column widths
    ws.column_dimensions["A"].width = 14  # Minggu Ke
    ws.column_dimensions["B"].width = 10  # Hari
    ws.column_dimensions["C"].width = 10  # Tanggal
    for ci in range(len(staff_columns)):
        col_letter = openpyxl.utils.get_column_letter(4 + ci)
        ws.column_dimensions[col_letter].width = max(14, len(staff_columns[ci]) + 4)

    # Freeze panes
    ws.freeze_panes = "D4"

    # Auto-filter
    ws.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(3 + len(staff_columns))}{3 + len(rows_data)}"

    # ── Write summary stats ──
    summary_row = 4 + len(rows_data) + 2
    ws.cell(row=summary_row, column=1, value="RINGKASAN").font = Font(
        name="Calibri", bold=True, size=11, color="1a5dc6"
    )

    total_staff = len(staff_columns)
    total_days = len(dates)
    ws.cell(row=summary_row + 1, column=1, value=f"Total Staff: {total_staff}").font = Font(
        name="Calibri", size=10
    )
    ws.cell(row=summary_row + 2, column=1, value=f"Total Hari: {total_days}").font = Font(
        name="Calibri", size=10
    )
    ws.cell(row=summary_row + 3, column=1, value=f"Total Jadwal: {len(schedules)}").font = Font(
        name="Calibri", size=10
    )

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─── Helpers ────────────────────────────────────────────────────


def _month_name(m: int) -> str:
    months = [
        "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember",
    ]
    return months[m] if 1 <= m <= 12 else f"Bulan {m}"


def _day_name_ina(d: int) -> str:
    days = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
    return days[d]


async def fetch_schedules(
    http_client: Any,
    supabase_url: str,
    service_key: str,
    month: int,
    year: int,
) -> list[dict]:
    """
    Fetch schedules from Supabase for a given month/year.
    """
    import httpx

    start_date = f"{year}-{month:02d}-01"
    end_date = f"{year}-{month:02d}-31"

    url = f"{supabase_url}/rest/v1/pos_attendance_schedules"
    params: dict[str, Any] = {
        "select": "userName,userSlug,roleName,date,shiftName,startTime,endTime,status",
        "date": f"gte.{start_date}",
        "order": "date.asc,userName.asc",
    }

    async with httpx.AsyncClient(timeout=30) as client:
        resp = await client.get(
            url,
            headers={
                "apikey": service_key,
                "Authorization": f"Bearer {service_key}",
            },
            params=params,
        )
        resp.raise_for_status()
        data = resp.json()
        # Filter client-side for the end date
        return [r for r in data if r.get("date", "") <= end_date]
